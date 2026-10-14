"""Build the project report (Markdown -> DOCX -> PDF) from the live knowledge base.

    python tools/build_report.py

The rule table, test tables and sample run are generated from the code, so the
report always matches the implementation.  Needs pandoc; PDF needs LibreOffice.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import sys
from contextlib import redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from heatwave_es import build_engine  # noqa: E402
from heatwave_es.explain import full_report  # noqa: E402
from heatwave_es.knowledge_base import RULES, SOURCES  # noqa: E402
from heatwave_es.scenarios import SCENARIOS  # noqa: E402

DOCS = ROOT / "docs"
BRAND_RED = "B7202E"   # Somaiya red (verify against the SVU brand manual)
BRAND_GREY = "3C3C3B"

REFERENCES = [
    "India Meteorological Department (IMD). *FAQ on Heat Wave* and heat-wave criteria; "
    "colour-coded impact-based warnings. https://mausam.imd.gov.in",
    "IMD. Warm night criteria, Climate Services, IMD Pune. https://imdpune.gov.in",
    "National Disaster Management Authority (NDMA). *Guidelines for Preparation of Action "
    "Plan - Prevention and Management of Heat Wave*, 2016 (revised 2019). https://ndma.gov.in",
    "Ministry of Health & Family Welfare, NCDC. *National Action Plan on Heat Related "
    "Illnesses*, National Programme on Climate Change and Human Health (NPCCHH), 2021.",
    "Rothfusz, L. P. (1990). *The Heat Index Equation*. NWS Southern Region Technical "
    "Attachment SR 90-23; NOAA WPC heat-index equation page.",
    "IMD / ICAR. Gramin Krishi Mausam Sewa - district agromet advisories for heat wave "
    "(crop and livestock management).",
    "Knowlton, K. et al. (2014). Development and implementation of South Asia's first "
    "heat-health action plan in Ahmedabad (India). *Int. J. Environ. Res. Public Health*, "
    "11(4), 3473-3492.",
    "Somaiya Vidyavihar University. AI Use Case KJS-CES-01: *Climate Intelligence for "
    "Heatwave Monitoring, Prediction, and Early Warning* (Dr. Radhika Kotecha).",
    "Giarratano, J. & Riley, G. *Expert Systems: Principles and Programming*, 4th ed., "
    "Thomson (forward chaining, Rete, conflict resolution, CLIPS).",
    "Russell, S. & Norvig, P. *Artificial Intelligence: A Modern Approach*, 4th ed. - "
    "Ch. 9, inference in first-order logic (forward/backward chaining).",
]


def esc(s: str) -> str:
    return str(s).replace("|", "\\|")


def rule_table() -> str:
    rows = ["| ID | Rule | IF | THEN | Sal. | Src. |", "|---|---|---|---|---|---|"]
    for r in RULES:
        rows.append(f"| {r.id} | {esc(r.name)} | {esc(r.if_text())} | "
                    f"{esc(r.then_text())} | {r.salience} | {r.source.upper()} |")
    return "\n".join(rows)


def fmt_inputs(d: dict) -> str:
    return ", ".join(f"{k}={'/'.join(v) if isinstance(v, list) else v}"
                     for k, v in d.items() if v != [])


def run_scenarios():
    eng = build_engine()
    out = []
    for sc in SCENARIOS:
        res = eng.run(sc["inputs"])
        ok = all(res.facts.get(k) == v for k, v in sc["expect"].items()) and \
            all(r in res.fired_ids() for r in sc.get("must_fire", [])) and \
            not any(r in res.fired_ids() for r in sc.get("must_not_fire", []))
        out.append((sc, res, ok))
    return eng, out


def test_table(results) -> str:
    rows = ["| TC | Scenario | Alert / outcome | Rules fired | Result |",
            "|---|---|---|---|---|"]
    for sc, res, ok in results:
        f = res.facts
        outcome = (f"ERROR (halt {res.halted_by})" if res.halted_by
                   else f"{f.get('alert')} / {f.get('hw_category')}")
        if f.get("emergency"):
            outcome += " + EMERGENCY"
        rows.append(f"| {sc['id']} | {esc(sc['title'])} | {outcome} | "
                    f"{len(res.trace)} | {'PASS' if ok else 'FAIL'} |")
    return "\n".join(rows)


def test_details(results) -> str:
    parts = []
    for sc, res, ok in results:
        f = res.facts
        concl = f.get("recommendation") or f"INPUT ERROR: {f.get('error')}"
        adv = sorted({a for a, _, _ in res.advice})
        parts.append(
            f"**{sc['id']}: {sc['title']}**\n\n"
            f"- *Inputs:* `{fmt_inputs(sc['inputs'])}`\n"
            f"- *Expected:* {', '.join(f'{k} = {v}' for k, v in sc['expect'].items())}\n"
            f"- *Rules fired (in order):* {' → '.join(res.fired_ids())}\n"
            f"- *Output:* {concl}"
            + (f"; advice for: {', '.join(adv)}" if adv else "") + "\n"
            f"- *Result:* **{'PASS' if ok else 'FAIL'}**\n")
    return "\n".join(parts)


def coverage_text() -> str:
    from heatwave_es.kb_check import coverage
    fired, _ = coverage()
    return f"{len(fired)}/{len(RULES)} (100 %)" if len(fired) == len(RULES) else f"{len(fired)}/{len(RULES)}"


def pytest_summary() -> str:
    p = subprocess.run([sys.executable, "-m", "pytest", "-q", str(ROOT / "tests")],
                       capture_output=True, text=True, cwd=ROOT)
    return p.stdout.strip().splitlines()[-1].strip("= ")


def make_reference_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"],
                   check=True)
    d = Document(path)
    red, grey = RGBColor.from_string(BRAND_RED), RGBColor.from_string(BRAND_GREY)
    by_name = {s.name: s for s in d.styles}
    for name in ["Normal", "Body Text", "First Paragraph", "Compact"]:
        if name in by_name:
            st = by_name[name]
            st.font.name, st.font.size = "Calibri", Pt(11)
    for name, size, colour in [("Title", 22, red), ("Subtitle", 13, grey),
                               ("Heading 1", 16, red), ("Heading 2", 13, grey),
                               ("Heading 3", 11.5, grey)]:
        st = by_name[name]
        st.font.name, st.font.size, st.font.color.rgb, st.font.bold = "Calibri", Pt(size), colour, True
    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = "Somaiya Vidyavihar University  |  K. J. Somaiya School of Engineering"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.color.rgb, run.font.size, run.font.bold = red, Pt(9), True
    fp = sec.footer.paragraphs[0]
    fp.text = "216U50E514 Artificial Intelligence  |  LAB CA Component 2  |  HeatGuard"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in fp.runs:
        run.font.color.rgb, run.font.size = grey, Pt(8)
    d.save(path)


def style_tables(path: Path) -> None:
    """Add borders, a red header row and small font to every table."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = Document(path)
    for t in d.tables:
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"), e.set(qn("w:sz"), "4"), e.set(qn("w:color"), "BFBFBF")
            borders.append(e)
        tblPr.append(borders)
        big = len(t.columns) >= 5
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                if i == 0 and t.rows[0].cells[0].text.strip():
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"), shd.set(qn("w:fill"), BRAND_RED)
                    cell._tc.get_or_add_tcPr().append(shd)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8 if big else 10)
                        if i == 0 and t.rows[0].cells[0].text.strip():
                            r.font.color.rgb, r.font.bold = RGBColor(255, 255, 255), True
    d.save(path)


def main() -> None:
    eng, results = run_scenarios()
    tc09 = next(r for s, r, _ in results if s["id"] == "TC09")
    md = (DOCS / "report_template.md").read_text(encoding="utf-8")
    fills = {
        "{{SOURCES}}": "\n".join(f"- **{k.upper()}** - {v}" for k, v in SOURCES.items()),
        "{{N_RULES}}": str(len(RULES)),
        "{{RULE_TABLE}}": rule_table(),
        "{{SAMPLE_RUN}}": full_report(eng, tc09, verbose=False),
        "{{N_TESTS}}": str(len(SCENARIOS)),
        "{{COVERAGE}}": coverage_text(),
        "{{PYTEST}}": pytest_summary(),
        "{{TEST_TABLE}}": test_table(results),
        "{{TEST_DETAILS}}": test_details(results),
        "{{REFERENCES}}": "\n".join(f"{i}. {r}" for i, r in enumerate(REFERENCES, 1)),
    }
    for k, v in fills.items():
        md = md.replace(k, v)
    out_md = DOCS / "HeatGuard_Report.md"
    out_md.write_text(md, encoding="utf-8")

    # standalone test report (D4) as well
    (DOCS / "Test_Case_Report.md").write_text(
        "# D4 - Test Case Report (auto-generated)\n\n"
        f"pytest: **{fills['{{PYTEST}}']}**\n\n{fills['{{TEST_TABLE}}']}\n\n"
        f"{fills['{{TEST_DETAILS}}']}", encoding="utf-8")

    if not shutil.which("pandoc"):
        print("pandoc not found - wrote Markdown only")
        return
    ref = DOCS / "reference.docx"
    make_reference_docx(ref)
    docx = DOCS / "HeatGuard_Report.docx"
    subprocess.run(["pandoc", str(out_md), "-o", str(docx), "--reference-doc", str(ref),
                    "--resource-path", str(DOCS)], check=True)
    style_tables(docx)
    ref.unlink()
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(DOCS),
                        str(docx)], check=True, capture_output=True)
    print("built:", ", ".join(p.name for p in DOCS.glob("HeatGuard_Report.*")))


if __name__ == "__main__":
    main()
